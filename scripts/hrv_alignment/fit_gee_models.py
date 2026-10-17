#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
 ACOUSTIC–HRV COUPLING ANALYSIS
=================================================

Implements the prespecified model using the packaged analysis input.

Primary question
----------------
Within the conservative fully support-matched O1 core, is within-work acoustic
variation associated with ln(RMSSD), after accounting for repeated cows and
broad linear temporal drift?

Primary family (4 tests)
------------------------
- roughness
- acoustic arousal (A_adj)
- BPM
- onset density

Primary model for each predictor
--------------------------------
Gaussian GEE, cow clusters, AR(1) working correlation, bias-reduced covariance:
    ln(RMSSD) ~ standardized_acoustic_predictor + centered_linear_time

Primary analytical core
-----------------------
Section-aware analytical core: 7 cows x 18 records (indices 3..20).
Each retained HRV estimate has a complete 300-s physiological support interval
inside O1 and acoustic predictors are exact section-timeweighted values across
that same support interval.

Multiplicity
------------
The 4 primary acoustic-predictor p-values are adjusted jointly using Holm.
No "best lag" search is performed.

Prespecified sensitivities
--------------------------
1. Exchangeable working correlation, retaining linear time.
2. AR(1), removing linear time.

Interpretive boundary
----------------------
This analysis tests acoustic–physiological association at the temporal
resolution of overlapping 300-s Kubios HRV estimates. It does NOT estimate a
physiological latency and does NOT establish acoustic causality.
"""

from __future__ import annotations

import hashlib
import argparse
import json
import math
import re
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm
from statsmodels.genmod.cov_struct import Autoregressive, Exchangeable
from statsmodels.genmod.families import Gaussian
from statsmodels.stats.multitest import multipletests


PACKAGE_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CORE = PACKAGE_ROOT / "data" / "derived" / "hrv_acoustic_alignment" / "O1_core_model_input.csv"
DEFAULT_OUTPUT = PACKAGE_ROOT / "data" / "derived" / "hrv_acoustic_alignment" / "model_results"

PREFIX = "acoustic_hrv_coupling"

PREDICTOR_SPECS = {
    "roughness": ["roughness"],
    "A_adj": ["a_adj", "a adj", "arousal"],
    "BPM": ["bpm"],
    "onset_density": ["onsets_per_s", "onsets per s", "onset_density", "onset density"],
}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower().replace("-", " ").replace("_", " "))


def find_core_input(path: Path) -> Path:
    resolved = path.resolve()
    if not resolved.is_file():
        raise FileNotFoundError(f"Core model-input CSV not found: {resolved}")
    return resolved


def find_col(df: pd.DataFrame, candidates, required=True):
    exact = {norm(c): c for c in df.columns}
    for cand in candidates:
        if norm(cand) in exact:
            return exact[norm(cand)]
    # substring fallback
    for c in df.columns:
        nc = norm(c)
        for cand in candidates:
            if norm(cand) in nc:
                return c
    if required:
        raise KeyError(f"Could not find {candidates}. Available columns: {list(df.columns)}")
    return None


def locate_predictor_columns(df: pd.DataFrame):
    tw_cols = [c for c in df.columns if "section_timeweighted" in str(c).lower()]
    if len(tw_cols) < 4:
        raise RuntimeError(
            f"Expected section-time-weighted acoustic predictors; found {len(tw_cols)}."
        )

    found = {}
    used = set()
    for label, toks in PREDICTOR_SPECS.items():
        hits = []
        for c in tw_cols:
            lc = str(c).lower()
            if any(tok.lower() in lc for tok in toks):
                hits.append(c)
        # Prefer the shortest/most exact-looking column.
        hits = sorted(hits, key=lambda c: (len(str(c)), str(c)))
        if not hits:
            raise RuntimeError(f"Could not identify frozen predictor '{label}' among: {tw_cols}")
        chosen = hits[0]
        if chosen in used:
            raise RuntimeError(f"Predictor column collision: {chosen}")
        used.add(chosen)
        found[label] = chosen
    return found


def validate_core(df, cow_col, rec_col, rmssd_col, pred_cols):
    df = df.copy()
    df[cow_col] = pd.to_numeric(df[cow_col], errors="coerce")
    df[rec_col] = pd.to_numeric(df[rec_col], errors="coerce")
    df[rmssd_col] = pd.to_numeric(df[rmssd_col], errors="coerce")
    for c in pred_cols.values():
        df[c] = pd.to_numeric(df[c], errors="coerce")

    if df[[cow_col, rec_col, rmssd_col] + list(pred_cols.values())].isna().any().any():
        bad = df[[cow_col, rec_col, rmssd_col] + list(pred_cols.values())].isna().sum()
        raise RuntimeError(f"Missing values in required primary variables:\n{bad}")

    df[cow_col] = df[cow_col].astype(int)
    df[rec_col] = df[rec_col].astype(int)

    counts = df.groupby(cow_col).size()
    if len(counts) != 7 or not (counts == 18).all():
        raise RuntimeError(f"Expected 7 cows x 18 core observations; got {counts.to_dict()}")
    recs = sorted(df[rec_col].unique().tolist())
    if recs != list(range(3, 21)):
        raise RuntimeError(f"Expected core record indices 3..20; got {recs}")
    if (df[rmssd_col] <= 0).any():
        raise RuntimeError("RMSSD contains non-positive values; ln(RMSSD) is not defined.")

    # Acoustic predictors must be identical across cows at a homologous record.
    for label, c in pred_cols.items():
        nun = df.groupby(rec_col)[c].nunique(dropna=False)
        if (nun != 1).any():
            raise RuntimeError(f"Frozen acoustic predictor '{label}' is not identical across cows by record.")
    return df


def fit_gee(d: pd.DataFrame, y_col: str, x_col: str, cow_col: str, rec_col: str,
            cov_kind: str, include_time: bool):
    cols = ["intercept", x_col]
    X = pd.DataFrame({"intercept": 1.0, x_col: d[x_col].astype(float)}, index=d.index)
    if include_time:
        X["time_centered_min"] = d["time_centered_min"].astype(float)
        cols.append("time_centered_min")

    if cov_kind == "AR1":
        cov_struct = Autoregressive(grid=True)
    elif cov_kind == "Exchangeable":
        cov_struct = Exchangeable()
    else:
        raise ValueError(cov_kind)

    model = sm.GEE(
        endog=d[y_col].astype(float),
        exog=X[cols],
        groups=d[cow_col].astype(int),
        time=d[[rec_col]].to_numpy(dtype=float),
        family=Gaussian(),
        cov_struct=cov_struct,
    )
    result = model.fit(cov_type="bias_reduced", maxiter=200)

    beta = float(result.params[x_col])
    se = float(result.bse[x_col])
    p = float(result.pvalues[x_col])
    ci = result.conf_int().loc[x_col]
    return {
        "beta_per_1SD_predictor": beta,
        "SE_bias_reduced": se,
        "CI95_low": float(ci.iloc[0]),
        "CI95_high": float(ci.iloc[1]),
        "P_raw": p,
        "percent_RMSSD_change_per_1SD": 100.0 * (math.exp(beta) - 1.0),
        "converged": bool(getattr(result, "converged", True)),
        "dep_params": float(np.asarray(result.cov_struct.dep_params).ravel()[0]) if result.cov_struct.dep_params is not None else np.nan,
        "n_rows": int(result.nobs),
        "n_cows": int(d[cow_col].nunique()),
    }


def style_xlsx(path: Path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    dark = "1F4E78"
    header_fill = PatternFill("solid", fgColor=dark)
    caution_fill = PatternFill("solid", fgColor="FCE4D6")
    teal_fill = PatternFill("solid", fgColor="DDEBF7")
    thin_gray = Side(style="thin", color="D9E1F2")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2" if ws.max_row > 1 else None
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            vals = [str(ws.cell(r, col_idx).value or "") for r in range(1, min(ws.max_row, 50) + 1)]
            width = min(max(max(map(len, vals), default=8) + 2, 10), 42)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if ws.title == "Primary_Results":
            for row in range(2, ws.max_row + 1):
                p = ws.cell(row, 10).value  # P_Holm
                if isinstance(p, (int, float)) and p < 0.05:
                    for cell in ws[row]:
                        cell.fill = teal_fill
        if ws.title == "README":
            for row in range(2, ws.max_row + 1):
                if "guardrail" in str(ws.cell(row, 1).value).lower() or "interpret" in str(ws.cell(row, 1).value).lower():
                    for cell in ws[row]:
                        cell.fill = caution_fill
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="Rebuild the prespecified O1 acoustic–HRV GEE models.")
    parser.add_argument("--core-input", type=Path, default=DEFAULT_CORE,
                        help="Compact 126-row input containing cow, record index, RMSSD, and four prespecified acoustic predictors.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    core_path = find_core_input(args.core_input)
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    print("=" * 100)
    print(" ACOUSTIC-HRV COUPLING ANALYSIS")
    print("=" * 100)
    print("[INPUT] Packaged model table:", core_path)
    print("[INFO] Primary core: records 3..20 only; expected 7 cows x 18 records")
    print("[INFO] Outcome: ln(RMSSD)")
    print("[INFO] Primary models: Gaussian GEE + AR(1) + bias-reduced covariance + centered linear time")
    print("[INFO] Four frozen predictors; Holm family = 4")
    print("[INFO] NO LAG SEARCH. NO BEST-LAG SELECTION. NO SI TESTING.")

    df = pd.read_csv(core_path)
    cow_col = find_col(df, ["cow", "vaca", "animal"])
    rec_col = find_col(df, ["o1_record_index_0based", "record_index", "record"])
    rmssd_col = find_col(df, ["rmssd"])
    pred_cols = locate_predictor_columns(df)
    df = validate_core(df, cow_col, rec_col, rmssd_col, pred_cols)

    df = df.sort_values([cow_col, rec_col]).reset_index(drop=True)
    df["log_RMSSD"] = np.log(df[rmssd_col].astype(float))
    df["time_centered_min"] = df[rec_col].astype(float) - df[rec_col].astype(float).mean()

    # Standardize predictors over the unique 18-point acoustic timeline, then map to all cows.
    standardization_rows = []
    for label, c in pred_cols.items():
        unique = df[[rec_col, c]].drop_duplicates(subset=[rec_col]).sort_values(rec_col)
        mu = float(unique[c].mean())
        sd = float(unique[c].std(ddof=1))
        if not np.isfinite(sd) or sd <= 0:
            raise RuntimeError(f"Predictor {label} has invalid SD={sd}")
        zcol = f"z_{label}"
        df[zcol] = (df[c].astype(float) - mu) / sd
        standardization_rows.append({
            "predictor": label,
            "source_column": c,
            "unique_timepoints": int(len(unique)),
            "mean_across_18_timepoints": mu,
            "sd_across_18_timepoints": sd,
            "min": float(unique[c].min()),
            "max": float(unique[c].max()),
        })

    primary_rows = []
    sensitivity_rows = []

    for label in ["roughness", "A_adj", "BPM", "onset_density"]:
        zcol = f"z_{label}"
        res = fit_gee(df, "log_RMSSD", zcol, cow_col, rec_col, "AR1", True)
        primary_rows.append({"predictor": label, "source_column": pred_cols[label], "covariance": "AR1", "linear_time": True, **res})

        res_ex = fit_gee(df, "log_RMSSD", zcol, cow_col, rec_col, "Exchangeable", True)
        sensitivity_rows.append({"sensitivity": "Exchangeable_with_linear_time", "predictor": label, **res_ex})

        res_nt = fit_gee(df, "log_RMSSD", zcol, cow_col, rec_col, "AR1", False)
        sensitivity_rows.append({"sensitivity": "AR1_without_linear_time", "predictor": label, **res_nt})

    primary = pd.DataFrame(primary_rows)
    reject, p_holm, _, _ = multipletests(primary["P_raw"].to_numpy(float), alpha=0.05, method="holm")
    primary["P_Holm"] = p_holm
    primary["Holm_reject_0.05"] = reject

    # Put adjusted p near raw p for readability.
    cols = [
        "predictor", "source_column", "covariance", "linear_time",
        "beta_per_1SD_predictor", "SE_bias_reduced", "CI95_low", "CI95_high",
        "P_raw", "P_Holm", "Holm_reject_0.05", "percent_RMSSD_change_per_1SD",
        "dep_params", "n_rows", "n_cows", "converged"
    ]
    primary = primary[cols]
    sens = pd.DataFrame(sensitivity_rows)

    readme = pd.DataFrame([
        ["Primary outcome", "ln(RMSSD) only"],
        ["Biological unit", "Cow; n=7"],
        ["Repeated observations", "18 fully support-matched Kubios estimates per cow (records 3..20); 126 rows total"],
        ["HRV geometry", "Each estimate summarizes a 300-s RR window; grid interval 60 s; consecutive estimates overlap 80%."],
        ["Acoustic predictors", "Roughness, A_adj, BPM and onset-density section-time-weighted values from the packaged input."],
        ["Primary model", "Separate Gaussian GEE per predictor; cow clusters; AR(1); bias-reduced covariance; centered linear time nuisance term."],
        ["Predictor scaling", "Each acoustic predictor standardized over the 18 unique core time positions before model fitting."],
        ["Multiplicity", "Four primary acoustic-predictor p-values adjusted jointly using Holm."],
        ["Sensitivities", "Exchangeable correlation with time; AR(1) without linear time. Sensitivities do not replace primary model by significance."],
        ["No lag search", "No nominal lag 0-3 family is searched or optimized in this analysis."],
        ["Interpretive guardrail", "Association at overlapping 300-s HRV support resolution; not physiological latency, causality, affective decoding, or proof of parasympathetic response to a specific acoustic feature."],
    ], columns=["Item", "Specification"])

    out_xlsx = output_dir / f"{PREFIX}.xlsx"
    out_primary = output_dir / f"{PREFIX}_primary_results.csv"
    out_sens = output_dir / f"{PREFIX}_sensitivities.csv"
    out_meta = output_dir / f"{PREFIX}_metadata.json"

    primary.to_csv(out_primary, index=False)
    sens.to_csv(out_sens, index=False)

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as w:
        readme.to_excel(w, sheet_name="README", index=False)
        primary.to_excel(w, sheet_name="Primary_Results", index=False)
        sens.to_excel(w, sheet_name="Sensitivities", index=False)
        pd.DataFrame(standardization_rows).to_excel(w, sheet_name="Predictor_Scaling", index=False)
        df[[cow_col, rec_col, rmssd_col, "log_RMSSD", "time_centered_min"] + list(pred_cols.values()) + [f"z_{x}" for x in pred_cols]].to_excel(w, sheet_name="Analysis_Data", index=False)

    style_xlsx(out_xlsx)

    metadata = {
        "build_status": "reproduced_from_packaged_model_input",
        "model_input": "data/derived/hrv_acoustic_alignment/" + core_path.name,
        "input_core_sha256": sha256_file(core_path),
        "primary_outcome": "ln(RMSSD)",
        "biological_n_cows": 7,
        "records_per_cow": 18,
        "record_indices": list(range(3, 21)),
        "n_repeated_rows": 126,
        "hrv_window_width_s": 300,
        "hrv_grid_interval_s": 60,
        "consecutive_window_overlap_fraction": 0.80,
        "predictor_columns": pred_cols,
        "primary_model": "Gaussian GEE; cow clusters; AR(1); bias-reduced covariance; centered linear time",
        "multiplicity": "Holm across 4 primary predictor tests",
        "sensitivities": ["Exchangeable with linear time", "AR1 without linear time"],
        "lag_search_performed": False,
        "best_lag_selected": False,
        "stress_index_tested": False,
        "interpretive_guardrail": "Association at overlapping 300-s HRV support resolution; not physiological latency or causality.",
        "outputs": [p.name for p in [out_xlsx, out_primary, out_sens]],
    }
    out_meta.write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")

    print("\nPRIMARY RESULTS — Holm family of 4")
    print("-" * 100)
    print(primary.to_string(index=False))

    print("\nPRESPECIFIED SENSITIVITIES")
    print("-" * 100)
    show = ["sensitivity", "predictor", "beta_per_1SD_predictor", "CI95_low", "CI95_high", "P_raw", "percent_RMSSD_change_per_1SD", "dep_params", "converged"]
    print(sens[show].to_string(index=False))

    print("\n[OUTPUTS]")
    for p in [out_xlsx, out_primary, out_sens, out_meta]:
        print(p)
    print("=" * 100)
    print("[COMPLETE] Acoustic-to-HRV models finished. No lag search was performed.")
    print("=" * 100)


if __name__ == "__main__":
    main()
